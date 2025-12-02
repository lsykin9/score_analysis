"""
生成成绩分析Excel模板
支持22列格式：姓名 + 总分(3列) + 6科×3列
"""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# 创建示例数据（22列格式）
data = {
    '姓名': ['张三', '李四', '王五'],
    '总分': [650, 620, 590],
    '总分年级排名': [15, 28, 45],
    '总分集团排名': [120, 150, 200],
    '语文': [125, 120, 115],
    '语文年级排名': [20, 35, 50],
    '语文集团排名': [100, 140, 180],
    '数学': [140, 130, 125],
    '数学年级排名': [10, 25, 40],
    '数学集团排名': [80, 130, 170],
    '英语': [135, 128, 120],
    '英语年级排名': [18, 30, 48],
    '英语集团排名': [110, 145, 190],
    '物理': [85, 80, 75],
    '物理年级排名': [12, 26, 42],
    '物理集团排名': [95, 135, 175],
    '化学': [90, 88, 82],
    '化学年级排名': [15, 28, 45],
    '化学集团排名': [105, 140, 185],
    '生物': [75, 74, 73],
    '生物年级排名': [20, 32, 46],
    '生物集团排名': [115, 148, 195]
}

df = pd.DataFrame(data)

# 保存为Excel
excel_path = '成绩分析模板.xlsx'
df.to_excel(excel_path, index=False, engine='openpyxl')

# 美化Excel
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# 设置样式
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 设置表头样式
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

# 设置数据区域样式和边框
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

# 调整列宽
ws.column_dimensions['A'].width = 12  # 姓名
for col in ['B', 'C', 'D']:  # 总分相关
    ws.column_dimensions[col].width = 14

# 各科列
for col_letter in ['E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V']:
    ws.column_dimensions[col_letter].width = 14

# 添加说明行
ws.insert_rows(1)
ws.merge_cells('A1:V1')
ws['A1'] = '说明：此模板为22列格式，包含姓名+总分(分数+年级排名+集团排名)+6科×3列。如无集团排名，可删除相关列。缺考学生对应科目填0。'
ws['A1'].font = Font(bold=True, color='FF0000', size=10)
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[1].height = 30

wb.save(excel_path)
print(f'✅ 模板已生成：{excel_path}')
print(f'📊 格式：22列（姓名 + 总分3列 + 6科×3列）')
print(f'👥 示例数据：3名学生')
print(f'📝 请根据实际情况修改数据后上传')
