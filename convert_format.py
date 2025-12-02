"""
将学生成绩.xlsx转换为22列标准格式
提取：姓名 + 总分(分数+年级排名+集团排名) + 6科×3列
"""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# 读取原始文件（使用多层表头）
df_raw = pd.read_excel('学生成绩.xlsx', header=[0, 1])

# 提取需要的数据
results = []
for idx, row in df_raw.iterrows():
    # 基本信息
    name = row[('姓名', 'Unnamed: 2_level_1')]
    
    # 跳过无效行
    if pd.isna(name):
        continue
    
    # 总分信息
    total_score = row[('总分', '分数')]
    total_school_rank = row[('总分', '校名')]  # 校名 = 年级排名
    total_group_rank = row[('总分', '级名')]   # 级名 = 集团排名
    
    # 6科信息
    # 语文
    chinese_score = row[('语文', '分数')]
    chinese_school_rank = row[('语文', '校名')]
    chinese_group_rank = row[('语文', '级名')]
    
    # 数学
    math_score = row[('数学', '分数')]
    math_school_rank = row[('数学', '校名')]
    math_group_rank = row[('数学', '级名')]
    
    # 英语
    english_score = row[('英语', '分数')]
    english_school_rank = row[('英语', '校名')]
    english_group_rank = row[('英语', '级名')]
    
    # 物理
    physics_score = row[('物理', '分数')]
    physics_school_rank = row[('物理', '校名')]
    physics_group_rank = row[('物理', '级名')]
    
    # 化学
    chemistry_score = row[('化学', '分数')]
    chemistry_school_rank = row[('化学', '校名')]
    chemistry_group_rank = row[('化学', '级名')]
    
    # 生物
    biology_score = row[('生物', '分数')]
    biology_school_rank = row[('生物', '校名')]
    biology_group_rank = row[('生物', '级名')]
    
    results.append({
        '姓名': name,
        '总分': total_score,
        '总分年级排名': total_school_rank,
        '总分集团排名': total_group_rank,
        '语文': chinese_score,
        '语文年级排名': chinese_school_rank,
        '语文集团排名': chinese_group_rank,
        '数学': math_score,
        '数学年级排名': math_school_rank,
        '数学集团排名': math_group_rank,
        '英语': english_score,
        '英语年级排名': english_school_rank,
        '英语集团排名': english_group_rank,
        '物理': physics_score,
        '物理年级排名': physics_school_rank,
        '物理集团排名': physics_group_rank,
        '化学': chemistry_score,
        '化学年级排名': chemistry_school_rank,
        '化学集团排名': chemistry_group_rank,
        '生物': biology_score,
        '生物年级排名': biology_school_rank,
        '生物集团排名': biology_group_rank
    })

# 创建DataFrame
df_converted = pd.DataFrame(results)

# 处理NaN值（缺考用0表示）
df_converted = df_converted.fillna(0)

# 保存为Excel
output_file = '学生成绩_转换后.xlsx'
df_converted.to_excel(output_file, index=False, engine='openpyxl')

# 美化Excel
wb = openpyxl.load_workbook(output_file)
ws = wb.active

# 设置样式
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 设置表头样式
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

# 设置数据区域样式
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

# 调整列宽
ws.column_dimensions['A'].width = 12  # 姓名
for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V']:
    ws.column_dimensions[col_letter].width = 14

wb.save(output_file)

print(f'✅ 转换完成！')
print(f'📁 输出文件：{output_file}')
print(f'📊 格式：22列（姓名 + 总分3列 + 6科×3列）')
print(f'👥 学生数量：{len(results)}')
print(f'\n前3名学生数据验证：')
print(df_converted[['姓名', '总分', '总分年级排名', '语文', '语文年级排名', '数学', '数学年级排名']].head(3))
print(f'\n说明：')
print(f'- 年级排名 = 原文件的"校名"列')
print(f'- 集团排名 = 原文件的"级名"列')
print(f'- 缺考数据已填充为0')
