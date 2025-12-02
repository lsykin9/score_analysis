"""
生成测试用的成绩总表 - 简化版本
包含4次考试数据，涵盖各种测试场景
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# 设置随机种子以确保可复现
np.random.seed(42)

# 读取原始数据
df_source = pd.read_excel('学生成绩_转换后.xlsx')
students = df_source['姓名'].tolist()
print(f"读取到 {len(students)} 个学生")
print(f"前6个学生: {students[:6]}")

# 科目列表
subjects = ['语文', '数学', '英语', '物理', '化学', '生物']

# 初始化DataFrame
df_final = pd.DataFrame({'姓名': students})

# 为每次考试生成数据
for exam_num in range(1, 5):
    print(f"\n正在生成第{exam_num}次考试数据...")
    
    # 为每个学生生成成绩
    all_scores = {}  # {科目: [学生1分数, 学生2分数, ...]}
    for subj in subjects:
        all_scores[subj] = []
    
    for i, student in enumerate(students):
        # === 第一次考试：基准 ===
        if exam_num == 1:
            scores = {
                '语文': np.random.randint(95, 125),
                '数学': np.random.randint(90, 130),
                '英语': np.random.randint(95, 125),
                '物理': np.random.randint(70, 95),
                '化学': np.random.randint(70, 95),
                '生物': np.random.randint(65, 90)
            }
            
        # === 第二次考试：引入缺考和偏科 ===
        elif exam_num == 2:
            if i == 0:  # 林子桓：数学缺考
                scores = {
                    '语文': np.random.randint(100, 120),
                    '数学': 0,  # 缺考
                    '英语': np.random.randint(100, 120),
                    '物理': np.random.randint(75, 90),
                    '化学': np.random.randint(75, 90),
                    '生物': np.random.randint(70, 85)
                }
                print(f"  学生{i+1}({student}): 数学缺考")
                
            elif i == 1:  # 张泽轩：英语缺考
                scores = {
                    '语文': np.random.randint(110, 125),
                    '数学': np.random.randint(110, 125),
                    '英语': 0,  # 缺考
                    '物理': np.random.randint(80, 92),
                    '化学': np.random.randint(80, 92),
                    '生物': np.random.randint(75, 88)
                }
                print(f"  学生{i+1}({student}): 英语缺考")
                
            elif i == 2:  # 李嘉然：偏科（数学强语文弱）
                scores = {
                    '语文': np.random.randint(85, 100),
                    '数学': np.random.randint(125, 135),
                    '英语': np.random.randint(105, 118),
                    '物理': np.random.randint(88, 95),
                    '化学': np.random.randint(88, 95),
                    '生物': np.random.randint(82, 90)
                }
                print(f"  学生{i+1}({student}): 偏科(数学强)")
                
            elif i == 3:  # 唐闻稻：不偏科
                avg = np.random.randint(105, 115)
                scores = {
                    '语文': avg + np.random.randint(-5, 5),
                    '数学': avg + np.random.randint(-5, 5),
                    '英语': avg + np.random.randint(-5, 5),
                    '物理': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '化学': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '生物': int((avg + np.random.randint(-5, 5)) * 0.72)
                }
                print(f"  学生{i+1}({student}): 不偏科")
                
            else:
                scores = {
                    '语文': np.random.randint(95, 125),
                    '数学': np.random.randint(90, 130),
                    '英语': np.random.randint(95, 125),
                    '物理': np.random.randint(70, 95),
                    '化学': np.random.randint(70, 95),
                    '生物': np.random.randint(65, 90)
                }
                
        # === 第三次考试：缺考后的表现 ===
        elif exam_num == 3:
            if i == 0:  # 林子桓：数学进步
                scores = {
                    '语文': np.random.randint(100, 120),
                    '数学': np.random.randint(125, 135),  # 高分
                    '英语': np.random.randint(100, 120),
                    '物理': np.random.randint(78, 92),
                    '化学': np.random.randint(78, 92),
                    '生物': np.random.randint(72, 87)
                }
                print(f"  学生{i+1}({student}): 数学进步到{scores['数学']}")
                
            elif i == 1:  # 张泽轩：英语退步
                scores = {
                    '语文': np.random.randint(110, 125),
                    '数学': np.random.randint(110, 125),
                    '英语': np.random.randint(80, 95),  # 低分
                    '物理': np.random.randint(80, 92),
                    '化学': np.random.randint(80, 92),
                    '生物': np.random.randint(75, 88)
                }
                print(f"  学生{i+1}({student}): 英语退步到{scores['英语']}")
                
            elif i == 2:  # 李嘉然：继续偏科
                scores = {
                    '语文': np.random.randint(82, 98),
                    '数学': np.random.randint(128, 138),
                    '英语': np.random.randint(103, 120),
                    '物理': np.random.randint(86, 95),
                    '化学': np.random.randint(86, 95),
                    '生物': np.random.randint(80, 90)
                }
                
            elif i == 3:  # 唐闻稻：继续均衡
                avg = np.random.randint(108, 118)
                scores = {
                    '语文': avg + np.random.randint(-5, 5),
                    '数学': avg + np.random.randint(-5, 5),
                    '英语': avg + np.random.randint(-5, 5),
                    '物理': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '化学': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '生物': int((avg + np.random.randint(-5, 5)) * 0.72)
                }
                
            elif i == 4:  # 庄曼丽：化学缺考
                scores = {
                    '语文': np.random.randint(105, 120),
                    '数学': np.random.randint(100, 120),
                    '英语': np.random.randint(105, 120),
                    '物理': np.random.randint(75, 90),
                    '化学': 0,  # 缺考
                    '生物': np.random.randint(70, 85)
                }
                print(f"  学生{i+1}({student}): 化学缺考")
                
            else:
                scores = {
                    '语文': np.random.randint(95, 125),
                    '数学': np.random.randint(90, 130),
                    '英语': np.random.randint(95, 125),
                    '物理': np.random.randint(70, 95),
                    '化学': np.random.randint(70, 95),
                    '生物': np.random.randint(65, 90)
                }
                
        # === 第四次考试：最终测试 ===
        elif exam_num == 4:
            if i == 0:  # 林子桓：继续保持
                scores = {
                    '语文': np.random.randint(105, 122),
                    '数学': np.random.randint(128, 138),
                    '英语': np.random.randint(105, 122),
                    '物理': np.random.randint(80, 94),
                    '化学': np.random.randint(80, 94),
                    '生物': np.random.randint(74, 89)
                }
                
            elif i == 1:  # 张泽轩：英语继续低迷
                scores = {
                    '语文': np.random.randint(112, 127),
                    '数学': np.random.randint(112, 127),
                    '英语': np.random.randint(75, 92),
                    '物理': np.random.randint(82, 94),
                    '化学': np.random.randint(82, 94),
                    '生物': np.random.randint(77, 90)
                }
                
            elif i == 2:  # 李嘉然：偏科更明显
                scores = {
                    '语文': np.random.randint(80, 95),
                    '数学': np.random.randint(130, 140),
                    '英语': np.random.randint(100, 118),
                    '物理': np.random.randint(84, 95),
                    '化学': np.random.randint(84, 95),
                    '生物': np.random.randint(78, 90)
                }
                
            elif i == 3:  # 唐闻稻：继续均衡
                avg = np.random.randint(110, 120)
                scores = {
                    '语文': avg + np.random.randint(-5, 5),
                    '数学': avg + np.random.randint(-5, 5),
                    '英语': avg + np.random.randint(-5, 5),
                    '物理': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '化学': int((avg + np.random.randint(-5, 5)) * 0.75),
                    '生物': int((avg + np.random.randint(-5, 5)) * 0.72)
                }
                
            elif i == 4:  # 庄曼丽：化学进步
                scores = {
                    '语文': np.random.randint(108, 122),
                    '数学': np.random.randint(103, 122),
                    '英语': np.random.randint(108, 122),
                    '物理': np.random.randint(78, 92),
                    '化学': np.random.randint(85, 95),  # 进步
                    '生物': np.random.randint(73, 87)
                }
                print(f"  学生{i+1}({student}): 化学进步到{scores['化学']}")
                
            elif i == 5:  # 薛毅：物理偏弱
                scores = {
                    '语文': np.random.randint(110, 125),
                    '数学': np.random.randint(110, 125),
                    '英语': np.random.randint(110, 125),
                    '物理': np.random.randint(55, 70),  # 弱项
                    '化学': np.random.randint(85, 95),
                    '生物': np.random.randint(80, 90)
                }
                print(f"  学生{i+1}({student}): 物理偏弱={scores['物理']}")
                
            else:
                scores = {
                    '语文': np.random.randint(95, 125),
                    '数学': np.random.randint(90, 130),
                    '英语': np.random.randint(95, 125),
                    '物理': np.random.randint(70, 95),
                    '化学': np.random.randint(70, 95),
                    '生物': np.random.randint(65, 90)
                }
        
        # 收集各科成绩
        for subj in subjects:
            all_scores[subj].append(scores[subj])
    
    # 计算总分
    total_scores = []
    for idx in range(len(students)):
        total = sum(all_scores[subj][idx] for subj in subjects)
        total_scores.append(total)
    
    # 添加总分和排名
    total_ranks = pd.Series(total_scores).rank(method='min', ascending=False).astype(int).tolist()
    group_ranks = [min(1500, max(1, rank + np.random.randint(-50, 100))) for rank in total_ranks]
    
    df_final[f'总分_考试{exam_num}'] = total_scores
    df_final[f'总分_年级排名_考试{exam_num}'] = total_ranks
    df_final[f'总分_集团排名_考试{exam_num}'] = group_ranks
    
    # 添加各科成绩和排名
    for subj in subjects:
        # 成绩
        df_final[f'{subj}_考试{exam_num}'] = all_scores[subj]
        
        # 年级排名（只对非0分数排名）
        subject_series = pd.Series(all_scores[subj])
        valid_mask = subject_series > 0
        ranks = [0] * len(students)
        if valid_mask.any():
            valid_ranks = subject_series[valid_mask].rank(method='min', ascending=False).astype(int)
            rank_idx = 0
            for idx in range(len(students)):
                if valid_mask.iloc[idx]:
                    ranks[idx] = valid_ranks.iloc[rank_idx]
                    rank_idx += 1
        df_final[f'{subj}_年级排名_考试{exam_num}'] = ranks
        
        # 集团排名
        group_subject_ranks = [min(1500, max(1, r + np.random.randint(-50, 100))) if r > 0 else 0 for r in ranks]
        df_final[f'{subj}_集团排名_考试{exam_num}'] = group_subject_ranks

# 保存到Excel
print(f"\n正在保存到Excel...")
wb = Workbook()
ws = wb.active
ws.title = "成绩总表"

# 写入数据
for r_idx, row in enumerate(dataframe_to_rows(df_final, index=False, header=True), 1):
    ws.append(row)

# 设置表头样式
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
header_alignment = Alignment(horizontal='center', vertical='center')

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment

# 设置边框
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border
        if cell.row > 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')

# 设置列宽
ws.column_dimensions['A'].width = 12
for col_idx in range(2, ws.max_column + 1):
    ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = 15

ws.freeze_panes = 'B2'

output_file = '测试成绩总表_4次考试.xlsx'
wb.save(output_file)

print(f"\n✅ 测试数据生成完成！")
print(f"文件: {output_file}")
print(f"学生数: {len(students)}")
print(f"考试数: 4")
print(f"总列数: {len(df_final.columns)}")
print(f"\n测试场景:")
print(f"  1. 林子桓: 第2次数学缺考 → 第3、4次进步")
print(f"  2. 张泽轩: 第2次英语缺考 → 第3、4次退步")
print(f"  3. 李嘉然: 偏科(数学强语文弱)")
print(f"  4. 唐闻稻: 不偏科(均衡发展)")
print(f"  5. 庄曼丽: 第3次化学缺考 → 第4次进步")
print(f"  6. 薛毅: 第4次物理明显偏弱")
